#!/usr/bin/env python3
"""Convierte la hoja «1. Fuentes» del Excel en catalogo/catalogo_completo.json, con las claves de la ficha del taller.

Uso (necesita openpyxl: pip install openpyxl):
    python3 catalogo/exportar_catalogo.py ruta/al/Diccionario_de_datos_TallerDatos.xlsx
El Excel es la fuente de verdad del catálogo; este JSON es su copia legible para el repositorio. Vuelvan a correrlo
cada vez que cambie el Excel. Incluye las fuentes que no entraron, con su estado y su nota: es parte de la entrega."""
import json, pathlib, sys
from openpyxl import load_workbook

ESTADO = {"Integrada al lago": "integrado", "Candidata": "candidato", "No responde": "caido",
          "No existe (la inventó la IA)": "inexistente", "Descartada": "excluido"}
ws = load_workbook(sys.argv[1], data_only=True)["1. Fuentes"]
fichas = []
for fila in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
    if not fila[0]:
        continue
    f = list(fila) + [None] * 16
    fichas.append({"id": f[0], "nombre": f[1], "entidad": f[2], "tipo_publicador": f[3], "url": f[4], "formato": f[5],
                   "como_se_obtiene": f[6], "licencia": f[7], "personas": f[8], "cobertura": f[9],
                   "vigencia": f[10], "frecuencia": f[11], "probado": str(f[12]) if f[12] else None,
                   "estado": ESTADO.get(f[13], f[13]), "hallada_por": f[14], "nota": f[15]})
destino = pathlib.Path(__file__).resolve().parent / "catalogo_completo.json"
destino.write_text(json.dumps(fichas, ensure_ascii=False, indent=1), encoding="utf-8")
cuenta = {}
for f in fichas:
    cuenta[f["estado"]] = cuenta.get(f["estado"], 0) + 1
print(f"{len(fichas)} fichas -> {destino.name}", cuenta)
